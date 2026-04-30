"""
Template contract extraction + validation.

Goal:
  Fail fast (or warn) when a docxtpl/Jinja2 template references variables that
  are not present in the runtime context, to avoid generating reports that look
  "successful" but contain missing content.

Scope (pragmatic, Word-template friendly):
  - Extract `{{ ... }}` variables (top-level + dotted paths).
  - Extract `{% for <var> in <list> %}` loop lists, and the fields referenced
    as `<var>.field` / `<var>['field']` inside the same table.

Notes:
  - This is a heuristic parser intended for docxtpl-style templates where loops
    are usually placed in table rows/cells.
  - It does not aim to fully parse Jinja2 expressions.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Set, Tuple

from docx import Document

_JINJA_VAR_RE = re.compile(r"\{\{\s*(?P<expr>.*?)\s*\}\}")
_JINJA_FOR_RE = re.compile(
    r"\{%\s*for\s+(?P<var>[a-zA-Z_][a-zA-Z0-9_]*)\s+in\s+"
    r"(?P<list>[a-zA-Z_][a-zA-Z0-9_]*)\s*%}"
)

# Best-effort: grab the leading "variable-ish" token from a Jinja expression.
# Supports dotted paths: foo.bar.baz
_LEADING_PATH_RE = re.compile(
    r"^\s*(?P<path>[a-zA-Z_][a-zA-Z0-9_]*(?:\.[a-zA-Z_][a-zA-Z0-9_]*)*)"
)

_JINJA_BUILTINS = {"loop", "cycler", "namespace", "range", "dict", "list", "lipsum"}


@dataclass(frozen=True)
class TemplateContract:
    template_path: str
    required_paths: Tuple[str, ...]
    required_lists: Tuple[str, ...]
    loop_row_fields: Dict[str, Tuple[str, ...]]  # list_name -> required row fields


@dataclass(frozen=True)
class ContractValidation:
    ok: bool
    missing_paths: Tuple[str, ...]
    missing_lists: Tuple[str, ...]
    missing_row_fields: Dict[str, Tuple[str, ...]]
    missing_row_examples: Dict[
        str, Dict[str, int]
    ]  # list_name -> field -> first missing row index (0-based)


def _iter_doc_text(doc: Document) -> Iterable[str]:
    for p in doc.paragraphs:
        if p.text:
            yield p.text

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if cell.text:
                    yield cell.text

    for sec in doc.sections:
        for p in sec.header.paragraphs:
            if p.text:
                yield p.text
        for p in sec.footer.paragraphs:
            if p.text:
                yield p.text


def _extract_loop_contract_from_table(table) -> List[Tuple[str, str, Set[str]]]:
    """Return (list_name, loop_var, required_fields) tuples from one table."""
    found: List[Tuple[str, str, Set[str]]] = []

    # Scan the whole table text, but keep it bounded to this table to avoid
    # cross-contamination.
    table_text_parts: List[str] = []
    for row in table.rows:
        for cell in row.cells:
            if cell.text:
                table_text_parts.append(cell.text)
    table_text = "\n".join(table_text_parts)

    for m in _JINJA_FOR_RE.finditer(table_text):
        loop_var = m.group("var")
        list_name = m.group("list")

        dot_re = re.compile(rf"\b{re.escape(loop_var)}\.([a-zA-Z_][a-zA-Z0-9_]*)\b")
        bracket_re = re.compile(rf"{re.escape(loop_var)}\[['\"]([^'\"]+)['\"]\]")

        fields: Set[str] = set()
        for fm in dot_re.finditer(table_text):
            fields.add(fm.group(1))
        for fm in bracket_re.finditer(table_text):
            fields.add(fm.group(1))

        found.append((list_name, loop_var, fields))

    return found


def extract_template_contract(template_path: str) -> TemplateContract:
    template_p = Path(template_path)
    if not template_p.exists():
        raise FileNotFoundError(f"Template docx not found: {template_p}")
    if template_p.suffix.lower() != ".docx":
        raise ValueError(f"Template must be a .docx file: {template_p}")

    doc = Document(str(template_p))

    loop_vars: Set[str] = set()
    required_lists: Set[str] = set()
    loop_row_fields: Dict[str, Set[str]] = {}

    for tbl in doc.tables:
        for list_name, loop_var, fields in _extract_loop_contract_from_table(tbl):
            loop_vars.add(loop_var)
            required_lists.add(list_name)
            loop_row_fields.setdefault(list_name, set()).update(fields)

    required_paths: Set[str] = set()
    for text in _iter_doc_text(doc):
        for m in _JINJA_VAR_RE.finditer(text):
            expr = m.group("expr") or ""
            # Strip filters: a|b|c -> a
            expr = expr.split("|", 1)[0]
            m2 = _LEADING_PATH_RE.match(expr)
            if not m2:
                continue
            path = m2.group("path")
            if not path:
                continue
            if path in _JINJA_BUILTINS:
                continue

            # Exclude loop variables (e.g. row.xxx) from top-level required paths.
            if path in loop_vars:
                continue
            if any(path.startswith(f"{lv}.") for lv in loop_vars):
                continue

            required_paths.add(path)

    # Sorting for stable output
    return TemplateContract(
        template_path=str(template_p),
        required_paths=tuple(sorted(required_paths)),
        required_lists=tuple(sorted(required_lists)),
        loop_row_fields={
            k: tuple(sorted(v)) for k, v in sorted(loop_row_fields.items())
        },
    )


def _get_by_path(obj: Any, path: str) -> bool:
    """Best-effort dotted-path resolver for dict-like contexts."""
    if not path:
        return False
    cur: Any = obj
    for part in path.split("."):
        if isinstance(cur, Mapping):
            if part not in cur:
                return False
            cur = cur[part]
            continue
        if hasattr(cur, part):
            cur = getattr(cur, part)
            continue
        return False
    return True


def validate_contract(
    contract: TemplateContract, *, context: Mapping[str, Any]
) -> ContractValidation:
    missing_paths: List[str] = []
    for p in contract.required_paths:
        if not _get_by_path(context, p):
            missing_paths.append(p)

    missing_lists: List[str] = []
    missing_row_fields: Dict[str, List[str]] = {}
    missing_row_examples: Dict[str, Dict[str, int]] = {}

    for list_name in contract.required_lists:
        if list_name not in context:
            missing_lists.append(list_name)
            continue

        value = context.get(list_name)
        if not isinstance(value, (list, tuple)):
            missing_lists.append(list_name)
            continue

        required_fields = set(contract.loop_row_fields.get(list_name, ()))
        if not required_fields or not value:
            continue

        # Strict: every row should provide every field referenced in the template.
        for field in sorted(required_fields):
            for idx, row in enumerate(value):
                ok = False
                if isinstance(row, Mapping):
                    ok = field in row
                elif hasattr(row, field):
                    ok = True

                if not ok:
                    missing_row_fields.setdefault(list_name, []).append(field)
                    missing_row_examples.setdefault(list_name, {})[field] = idx
                    break

    # Deduplicate and sort
    missing_row_fields_sorted: Dict[str, Tuple[str, ...]] = {
        k: tuple(sorted(set(v))) for k, v in missing_row_fields.items()
    }

    ok = (not missing_paths) and (not missing_lists) and (not missing_row_fields_sorted)

    return ContractValidation(
        ok=ok,
        missing_paths=tuple(sorted(set(missing_paths))),
        missing_lists=tuple(sorted(set(missing_lists))),
        missing_row_fields=missing_row_fields_sorted,
        missing_row_examples=missing_row_examples,
    )
