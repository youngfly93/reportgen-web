# ruff: noqa: E402, I001

import sys
from datetime import date
from pathlib import Path

import pytest
import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm


ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from reportgen.core.batch_runner import (
    BatchValidateOptions,
    _expected_tables_from_excel,
)
from reportgen.core.excel_reader import ExcelReader
from reportgen.core.field_mapper import FieldMapper
from reportgen.core.project_detector import ProjectDetector
from reportgen.core.report_generator import ReportGenerator
from reportgen.core.template_bridge_358 import (
    build_tmb_summary,
    build_variants_for_template,
    enhance_report_data,
    load_panel_config,
)
from reportgen.core.template_renderer import TemplateRenderer
from reportgen.core.validation import validate_excel_data_common
from reportgen.models.excel_data import ExcelDataSource
from reportgen.models.report_data import ReportData


def _excel(
    tmp_path: Path,
    *,
    single_values=None,
    variations=None,
    tables=None,
) -> ExcelDataSource:
    path = tmp_path / "unknown.xlsx"
    path.write_bytes(b"placeholder")
    table_data = dict(tables or {})
    if variations is not None:
        table_data["Variations"] = variations
    return ExcelDataSource(
        file_path=str(path),
        single_values=single_values or {},
        table_data=table_data,
        sheet_names=list(table_data),
        metadata={},
    )


def test_reviewed_override_does_not_create_absent_erbb2_tip(tmp_path):
    excel_data = _excel(tmp_path, variations=[])
    report_data = enhance_report_data(
        ReportData(),
        excel_data,
        base_path=str(ROOT),
    )

    tips = report_data.get_table("targeted_drug_tips")
    assert all(row.get("gene") != "ERBB2" for row in tips)
    assert report_data.get_field("msi_status") == "未检测"
    assert report_data.get_field("tmb_status") == "未检测"


def test_reviewed_override_applies_when_erbb2_variant_is_detected(tmp_path):
    excel_data = _excel(
        tmp_path,
        variations=[
            {
                "Gene_Symbol": "ERBB2",
                "Transcript": "NM_004448.4",
                "Chr": "17",
                "ExIn_ID": "EX17",
                "cHGVS": "c.1979G>A",
                "pHGVS_S": "p.G660D",
                "Freq(%)": 12.3,
                "ExistInsmall358": 1,
                "ExistIn552": 1,
                "CLNSIG": "Pathogenic",
            }
        ],
    )
    report_data = enhance_report_data(
        ReportData(),
        excel_data,
        base_path=str(ROOT),
    )

    tips = report_data.get_table("targeted_drug_tips")
    erbb2 = [row for row in tips if row.get("gene") == "ERBB2"]
    assert erbb2
    assert "曲妥珠单抗" in erbb2[0].get("benefit_drugs", "")


def test_nccn_mutation_rows_do_not_include_cnv_or_fusion(tmp_path):
    excel_data = _excel(
        tmp_path,
        variations=[
            {
                "Gene_Symbol": "ERBB2",
                "Transcript": "NM_004448.4",
                "Chr": "17",
                "ExIn_ID": "EX20",
                "cHGVS": "c.2324_2325ins12",
                "pHGVS_S": "p.Y772_A775dup",
                "Freq(%)": 12.3,
                "ExistInsmall358": 1,
                "ExistIn552": "Ⅱ类",
                "CLNSIG": "Pathogenic",
            }
        ],
        tables={
            "Cnv": [
                {"Gene": "ERBB2", "Status": "扩增"},
                {"Gene": "EGFR", "Status": "扩增"},
            ],
            "Fusion": [
                {"Gene1": "BICC1", "Gene2": "FGFR2"},
            ],
        },
    )
    report_data = enhance_report_data(
        ReportData(),
        excel_data,
        base_path=str(ROOT),
    )

    assert report_data.get_field("nccn_ERBB2_MUT") == (
        "c.2324_2325ins12，p.Y772_A775dup"
    )
    assert report_data.get_field("nccn_ERBB2_AMP") == "CNV:扩增"
    assert report_data.get_field("nccn_FGFR123_MUT") == "未检出"
    assert report_data.get_field("nccn_FGFR123_FUSION") == "融合:BICC1-FGFR2"
    assert report_data.get_field("imm_hyper_EGFR_AMP") == "CNV:扩增"


def test_field_mapper_adds_legacy_fusion_aliases(tmp_path):
    excel_data = _excel(
        tmp_path,
        tables={
            "Fusion": [
                {
                    "Gene1": "EML4",
                    "Chr1": "chr2",
                    "Pos1": 42491832,
                    "Gene2": "ALK",
                    "Chr2": "chr2",
                    "Pos2": 29446394,
                    "Sv_type": "fusion",
                }
            ]
        },
    )

    report_data = FieldMapper(
        config_dir=str(ROOT / "config"), log_level="ERROR"
    ).map(excel_data)

    rows = report_data.get_table("fusion")
    assert rows
    assert "#Est_Type" in rows[0]
    assert rows[0]["#Est_Type"] == "fusion"
    assert "Freq1" in rows[0]


def test_missing_clnsig_is_not_defaulted_to_pathogenic(tmp_path):
    excel_data = _excel(
        tmp_path,
        variations=[
            {
                "Gene_Symbol": "TP53",
                "Transcript": "NM_000546.6",
                "Chr": "17",
                "ExIn_ID": "EX8",
                "cHGVS": "c.817C>T",
                "pHGVS_S": "p.R273C",
                "Freq(%)": 33.0,
                "ExistInsmall358": 1,
                "ExistIn552": 1,
                "CLNSIG": "*",
            }
        ],
    )
    variants = build_variants_for_template(
        excel_data,
        filter_class_i_ii_only=False,
        important_genes_only=False,
        panel_config=load_panel_config(base_path=str(ROOT)),
    )

    assert variants
    assert variants[0]["clinical_significance"] == "临床意义未明"


def test_tmb_missing_is_explicit_not_low(tmp_path):
    summary = build_tmb_summary(_excel(tmp_path))
    assert summary["tmb_value"] == "未检测"
    assert summary["tmb_status"] == "未检测"
    assert summary["tmb_level_cn"] == "未检测"


def test_tmb_invalid_is_explicit_format_error(tmp_path):
    summary = build_tmb_summary(_excel(tmp_path, single_values={"TMB": "abc"}))
    assert summary["tmb_value"] == "未检测（格式错误）"
    assert summary["tmb_status"] == "未检测"
    assert summary["tmb_summary"] == "未检测（格式错误）"


def test_field_mapper_invalid_tmb_uses_same_format_error(tmp_path):
    report_data = FieldMapper(
        config_dir=str(ROOT / "config"), log_level="ERROR"
    ).map(_excel(tmp_path, single_values={"TMB": "abc"}))

    assert report_data.get_field("tmb_value") == "未检测（格式错误）"
    assert report_data.get_field("tmb_status") == "未检测"
    assert report_data.get_field("tmb_summary") == "未检测（格式错误）"


def test_field_mapper_valid_tmb_overrides_default_status(tmp_path):
    report_data = FieldMapper(
        config_dir=str(ROOT / "config"), log_level="ERROR"
    ).map(_excel(tmp_path, single_values={"TMB": 7.10499382716049, "MSI状态": "MSI-H"}))

    assert report_data.get_field("tmb_value") == "7.1"
    assert report_data.get_field("tmb") == "7.1 mutations/Mb"
    assert report_data.get_field("tmb_status") == "L"
    assert report_data.get_field("tmb_level_cn") == "低"
    assert "TMB-L" in report_data.get_field("tmb_summary")
    assert report_data.get_field("msi_status") == "MSI-H"
    assert "MSI-H" in report_data.get_field("immuno_tips")
    assert "TMB-H的肿瘤" not in report_data.get_field("immuno_tips")
    assert "TMB水平较低" in report_data.get_field("tmb_detail_sentence")
    assert "MSI-H" in report_data.get_field("msi_detail_sentence")
    assert "MSI-H" in report_data.get_field("msi_tips")


def test_field_mapper_dynamic_tmb_msi_narratives_match_mss_low_tmb(tmp_path):
    report_data = FieldMapper(
        config_dir=str(ROOT / "config"), log_level="ERROR"
    ).map(_excel(tmp_path, single_values={"TMB": 7.74681481481482, "MSI状态": "MSS"}))

    assert report_data.get_field("tmb_value") == "7.7"
    assert report_data.get_field("tmb_status") == "L"
    assert "7.7mutations/Mb" in report_data.get_field("tmb_detail_sentence")
    assert "TMB水平较低" in report_data.get_field("tmb_detail_sentence")
    assert "TMB-H参考阈值" in report_data.get_field("tmb_detail_interpretation")
    assert report_data.get_field("tmb_drug_note") == ""
    assert "微卫星稳定（MSS）型" in report_data.get_field("msi_detail_sentence")
    assert "未提示MSI-H/dMMR" in report_data.get_field("msi_tips")


def test_case2_fixture_tmb_msi_mapping_if_available():
    fixture = ROOT / "output/xlsx_for_win/case2_highTMB_MSIH_MLB0002.result.xlsx"
    if not fixture.exists():
        pytest.skip("case2 fixture is not available in this checkout")

    excel_data = ExcelReader(config_dir=str(ROOT / "config"), log_level="ERROR").read(
        str(fixture), include_tables=True
    )
    report_data = FieldMapper(config_dir=str(ROOT / "config"), log_level="ERROR").map(
        excel_data
    )

    assert excel_data.single_values["TMB"] == pytest.approx(7.10499382716049)
    assert excel_data.single_values["MSI状态"] == "MSI-H"
    assert report_data.get_field("tmb_value") == "7.1"
    assert report_data.get_field("tmb_status") == "L"
    assert report_data.get_field("msi_status") == "MSI-H"


def test_field_mapper_updates_msi_status_cn_from_mss(tmp_path):
    report_data = FieldMapper(
        config_dir=str(ROOT / "config"), log_level="ERROR"
    ).map(_excel(tmp_path, single_values={"MSI状态": "MSS"}))

    assert report_data.get_field("msi_status") == "MSS"
    assert report_data.get_field("msi_status_cn") == "微卫星稳定型，MSS"


def test_missing_report_date_is_not_backfilled_to_today():
    report_data = ReportData()
    generator = ReportGenerator(config_dir=str(ROOT / "config"), log_level="ERROR")

    generator._mark_missing_report_date(report_data)

    assert report_data.get_field("report_date") == "未填写"
    assert "缺失必填字段: report_date" in report_data.validation_errors


def test_common_validation_warns_without_today_backfill(tmp_path):
    warnings = validate_excel_data_common(_excel(tmp_path), today=date(2026, 4, 19))

    report_date_warnings = [w for w in warnings if w.get("field") == "report_date"]
    assert report_date_warnings
    assert "不会自动回填今天" in report_date_warnings[0]["message"]


def test_detector_does_not_use_panel_column_numbers_as_crc_signal(tmp_path):
    excel_data = _excel(
        tmp_path,
        single_values={"TMB": 3.58},
        variations=[{"ExistInsmall358": 1, "Gene_Symbol": "TP53", "cHGVS": "c.1A>T"}],
    )
    detector = ProjectDetector(config_dir=str(ROOT / "config"), log_level="ERROR")
    result = detector.detect(str(Path(excel_data.file_path)), excel_data=excel_data)

    assert result["project_type"] != "crc_358_msi"


def test_detector_uses_trusted_filename_project_tokens(tmp_path):
    detector = ProjectDetector(config_dir=str(ROOT / "config"), log_level="ERROR")
    excel_data = _excel(tmp_path)

    path_301 = tmp_path / "_MLS2600000001_结直肠癌301基因+MSI_终版.xlsx"
    path_358 = tmp_path / "_MLS2600000002_结直肠癌358基因+MSI_终版.xlsx"
    path_301.write_bytes(b"placeholder")
    path_358.write_bytes(b"placeholder")

    assert detector.detect(str(path_301), excel_data=excel_data)["project_type"] == "crc_301_msi"
    assert detector.detect(str(path_358), excel_data=excel_data)["project_type"] == "crc_358_msi"


def test_msi_percentage_label_conflict_is_warned(tmp_path):
    from app.services.reportgen_bridge import ReportGenBridge

    excel_data = _excel(
        tmp_path,
        single_values={"MSI状态": "MSS", "MSI百分比": 40.0},
    )
    bridge = ReportGenBridge(
        config_dir=str(ROOT / "config"),
        template_dir=str(ROOT / "templates"),
    )

    warnings = bridge.validate_excel_data(excel_data)
    assert any(w.get("field") == "msi_status" and w.get("level") == "warning" for w in warnings)


def test_batch_options_accept_forced_project_type():
    opts = BatchValidateOptions(
        inputs=["dummy.xlsx"],
        project_type="crc_358_msi",
        project_name="结直肠癌358基因+MSI",
    )
    assert opts.project_type == "crc_358_msi"


def test_hla_expected_table_matches_default_hidden_policy(tmp_path):
    excel_data = _excel(tmp_path, tables={"HLA": [{"Locus": "HLA-A", "Type1": "01:01"}]})

    assert _expected_tables_from_excel(excel_data, show_hla_table=False)["hla"] is False
    assert _expected_tables_from_excel(excel_data, show_hla_table=True)["hla"] is True


def test_consultation_phone_is_enabled_in_config():
    settings = yaml.safe_load((ROOT / "config" / "settings.yaml").read_text(encoding="utf-8"))
    assert (settings.get("report_content") or {}).get("consultation_phone") == "022-87190699"


def test_patient_info_table_uses_lz_project_code_and_removes_qc_rows(tmp_path):
    docx_path = tmp_path / "patient_info_table.docx"
    doc = Document()
    table = doc.add_table(rows=5, cols=4)
    rows = [
        ("姓名：", "陈三", "样本类型：", "组织"),
        ("性别：", "男", "取材手段：", "-"),
        ("临床诊断：", "结直肠癌", "项目编码：", "MLS2601287001"),
        ("送检医院：", "某某医院", "送检科室：", "肿瘤科"),
        ("病理号：", "-", "采集日期：", "2026-03-05"),
    ]
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.save(docx_path)

    TemplateRenderer(log_level="ERROR")._apply_report_content_fixes(
        str(docx_path),
        {"report_number": "MLJY-LZ260525", "sample_id": "MLS2601287001"},
    )

    table = Document(docx_path).tables[0]
    row_texts = [" ".join(cell.text for cell in row.cells) for row in table.rows]
    all_text = "\n".join(row_texts)
    assert "项目编码： LZ260525" in all_text
    assert "MLS2601287001" not in all_text
    for removed in ("送检医院", "送检科室", "病理号", "采集日期"):
        assert removed not in all_text


def test_targeted_drug_tips_table_style_is_normalized(tmp_path):
    docx_path = tmp_path / "targeted_drug_tips.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    headers = [
        "基因",
        "突变位点",
        "潜在获益靶向药物\n（证据等级）",
        "可能耐药或慎重药物\n（证据等级）",
    ]
    values = [
        "KRAS",
        "c.38G>A,\np.G13D",
        "西妥昔单抗（A）",
        "依维莫司（C）",
    ]
    for idx, value in enumerate(headers):
        table.rows[0].cells[idx].text = value
    for idx, value in enumerate(values):
        table.rows[1].cells[idx].text = value
    doc.save(docx_path)

    TemplateRenderer(log_level="ERROR")._normalize_targeted_drug_tips_table(str(docx_path))

    table = Document(docx_path).tables[0]
    grid_widths = [col.get(qn("w:w")) for col in table._tbl.tblGrid.gridCol_lst]
    assert grid_widths == ["1170", "1758", "3530", "1837"]

    header_cell = table.rows[0].cells[0]
    body_gene_cell = table.rows[1].cells[0]
    body_site_cell = table.rows[1].cells[1]
    header_shd = header_cell._tc.tcPr.find(qn("w:shd"))
    body_shd = body_gene_cell._tc.tcPr.find(qn("w:shd"))
    assert header_shd.get(qn("w:fill")) == "00C4D8"
    assert body_shd.get(qn("w:fill")) == "FFFFFF"

    header_run = header_cell.paragraphs[0].runs[0]
    gene_run = body_gene_cell.paragraphs[0].runs[0]
    site_run = body_site_cell.paragraphs[0].runs[0]
    assert header_run.font.bold is True
    assert header_run.font.size.pt == 9
    assert str(header_run.font.color.rgb) == "FFFFFF"
    assert gene_run.font.size.pt == 9
    assert str(gene_run.font.color.rgb) == "0000FF"
    assert gene_run.font.underline is True
    assert site_run.font.underline is False
    assert body_site_cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_reviewed_result_table_style_expands_without_touching_qa_tables(tmp_path):
    docx_path = tmp_path / "reviewed_result_tables.docx"
    doc = Document()

    biomarkers = doc.add_table(rows=2, cols=3)
    for idx, value in enumerate(
        ["TMB/MSI/其它生物标志物检测结果", "TMB/MSI/其它生物标志物检测结果", "用药提示"]
    ):
        biomarkers.rows[0].cells[idx].text = value
    for idx, value in enumerate(["肿瘤突变负荷（TMB）", "18.8", "用药提示文本"]):
        biomarkers.rows[1].cells[idx].text = value

    drugs = doc.add_table(rows=2, cols=3)
    for idx, value in enumerate(["药物名称", "相关基因", "药物适应情况"]):
        drugs.rows[0].cells[idx].text = value
    for idx, value in enumerate(["瑞戈非尼", "VEGFR", "适应情况文本"]):
        drugs.rows[1].cells[idx].text = value

    immune = doc.add_table(rows=2, cols=3)
    for idx, value in enumerate(["基因", "检测结果", "临床解读"]):
        immune.rows[0].cells[idx].text = value
    for idx, value in enumerate(["MLH1", "未检出有害变异", "临床解读文本"]):
        immune.rows[1].cells[idx].text = value

    qa = doc.add_table(rows=2, cols=2)
    qa.rows[0].cells[0].text = "问题1"
    qa.rows[0].cells[1].text = "肿瘤患者为什么要进行基因检测？"
    qa.rows[1].cells[0].text = ""
    qa.rows[1].cells[1].text = "问答正文"
    doc.save(docx_path)

    TemplateRenderer(log_level="ERROR")._normalize_reviewed_result_tables(str(docx_path))

    doc = Document(docx_path)
    for idx in (0, 1, 2):
        table = doc.tables[idx]
        header_shd = table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd"))
        body_shd = table.rows[1].cells[0]._tc.tcPr.find(qn("w:shd"))
        assert header_shd.get(qn("w:fill")) == "00C4D8"
        assert body_shd.get(qn("w:fill")) == "FFFFFF"
        assert table.rows[0].cells[0].paragraphs[0].runs[0].font.bold is True

    qa_header_shd = doc.tables[3].rows[0].cells[0]._tc.tcPr.find(qn("w:shd"))
    qa_body_shd = doc.tables[3].rows[1].cells[0]._tc.tcPr.find(qn("w:shd"))
    assert qa_header_shd is None
    assert qa_body_shd is None


def test_signature_placeholder_is_removed_without_image(tmp_path):
    docx_path = tmp_path / "signature.docx"
    doc = Document()
    doc.add_paragraph("签名：__SIG_IMG__")
    doc.save(docx_path)

    TemplateRenderer(log_level="ERROR")._render_signature_placeholder(str(docx_path), {})

    assert "__SIG_IMG__" not in "\n".join(p.text for p in Document(docx_path).paragraphs)


def test_signature_layout_moves_report_date_to_separate_line(tmp_path):
    from shutil import copyfile

    source = ROOT / "templates/aligned_template_with_cnv_fusion_hla_FIXED.docx"
    docx_path = tmp_path / "signature_layout.docx"
    copyfile(source, docx_path)

    renderer = TemplateRenderer(log_level="ERROR")
    renderer._render_signature_placeholder(str(docx_path), {})
    renderer._apply_report_content_fixes(str(docx_path), {"report_date": "2026-04-26"})
    renderer._normalize_signature_layout(str(docx_path), {"report_date": "2026-04-26"})

    paragraphs = [p.text.strip() for p in Document(docx_path).paragraphs if p.text.strip()]
    signature_lines = [p for p in paragraphs if p.startswith("检测者：")]
    assert signature_lines
    assert all("报告日期" not in p for p in signature_lines)
    assert "报告日期：2026.04.26" in paragraphs


def test_template_tmb_msi_patient_narratives_are_dynamic():
    template = ROOT / "templates/aligned_template_with_cnv_fusion_hla_FIXED.docx"
    doc = Document(template)
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    text = "\n".join(texts)

    assert "13.5mutations/Mb" not in text
    assert "该肿瘤样本为 微卫星稳定（MSS）型" not in text
    assert "MSI-H的实体瘤通常具有免疫原性" not in text
    for token in [
        "{{ tmb_detail_sentence }}",
        "{{ tmb_detail_interpretation }}",
        "{{ tmb_drug_note }}",
        "{{ msi_detail_sentence }}",
        "{{ msi_detail_interpretation }}",
        "{{ msi_tips }}",
    ]:
        assert token in text


def test_patient_letter_is_native_docx_text_not_legacy_textbox():
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET

    template = ROOT / "templates/aligned_template_with_cnv_fusion_hla_FIXED.docx"
    doc = Document(template)
    text_parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.append(cell.text)
    text = "\n".join(text_parts)

    assert "致您的一封信" in text
    assert "尊敬的 {{ patient_name }} {{ patient_salutation }}：" in text
    assert "现代医学已经证明" in text

    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(template) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    textbox_text = "\n".join(
        "".join(t.text or "" for t in textbox.findall(".//w:t", ns))
        for textbox in root.findall(".//w:txbxContent", ns)
    )

    assert "现代医学已经证明" not in textbox_text
    assert "尊敬的" not in textbox_text


def test_patient_letter_body_block_is_not_pushed_to_right():
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET

    template = ROOT / "templates/aligned_template_with_cnv_fusion_hla_FIXED.docx"
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(template) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))

    matching_tables = []
    for table in root.findall(".//w:tbl", ns):
        text = "".join(t.text or "" for t in table.findall(".//w:t", ns))
        if "现代医学已经证明" in text:
            matching_tables.append(table)

    assert len(matching_tables) == 1
    widths = [
        int(tc_w.get(f"{{{ns['w']}}}w") or "0")
        for tc_w in matching_tables[0].findall(".//w:tcPr/w:tcW", ns)[:2]
    ]
    grid_widths = [
        int(grid_col.get(f"{{{ns['w']}}}w") or "0")
        for grid_col in matching_tables[0].findall("./w:tblGrid/w:gridCol", ns)[:2]
    ]
    assert widths[0] <= 800
    assert widths[1] >= 8000
    assert grid_widths[0] <= 800
    assert grid_widths[1] >= 8000


def test_patient_letter_salutation_follows_gender():
    generator = ReportGenerator(config_dir=str(ROOT / "config"), log_level="ERROR")
    report_data = ReportData()

    report_data.set_field("gender", "女")
    generator._set_patient_salutation(report_data)
    assert report_data.get_field("patient_salutation") == "女士"

    report_data.set_field("gender", "男")
    generator._set_patient_salutation(report_data)
    assert report_data.get_field("patient_salutation") == "先生"


def test_toc_decoration_line_is_moved_left_and_up(tmp_path):
    from shutil import copyfile
    from zipfile import ZipFile
    import xml.etree.ElementTree as ET

    source = ROOT / "templates/aligned_template_with_cnv_fusion_hla_FIXED.docx"
    docx_path = tmp_path / "toc_layout.docx"
    copyfile(source, docx_path)

    TemplateRenderer(log_level="ERROR")._normalize_toc_decoration_layout(str(docx_path))

    ns_w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    ns_wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    w_p = f"{{{ns_w}}}p"
    w_t = f"{{{ns_w}}}t"
    wp_anchor = f"{{{ns_wp}}}anchor"
    wp_doc_pr = f"{{{ns_wp}}}docPr"
    wp_extent = f"{{{ns_wp}}}extent"
    wp_position_h = f"{{{ns_wp}}}positionH"
    wp_position_v = f"{{{ns_wp}}}positionV"
    wp_pos_offset = f"{{{ns_wp}}}posOffset"

    with ZipFile(docx_path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))

    line_offsets = []
    circle_offsets = []
    for para in root.iter(w_p):
        text = "".join(t.text or "" for t in para.iter(w_t)).replace(" ", "")
        if "目录" not in text:
            continue
        for anchor in para.iter(wp_anchor):
            doc_pr = anchor.find(wp_doc_pr)
            extent = anchor.find(wp_extent)
            if doc_pr is None or extent is None:
                continue
            name = doc_pr.get("name", "")
            cx = int(extent.get("cx") or "0")
            cy = int(extent.get("cy") or "0")
            pos_h = anchor.find(wp_position_h)
            pos_v = anchor.find(wp_position_v)
            if pos_h is None or pos_v is None:
                continue
            x = int((pos_h.find(wp_pos_offset).text or "0"))
            y = int((pos_v.find(wp_pos_offset).text or "0"))
            if "直接连接符" in name and cx <= 100000 and cy >= 2000000:
                line_offsets.append((x, y))
            elif "椭圆" in name and cx <= 200000 and cy <= 200000:
                circle_offsets.append((x, y))

    assert line_offsets
    assert circle_offsets
    assert all(x == 127000 and y == 533400 for x, y in line_offsets)
    assert all(x == 92710 and y == 457200 for x, y in circle_offsets)


def test_section_layouts_are_normalized(tmp_path):
    docx_path = tmp_path / "sections.docx"
    doc = Document()
    doc.add_paragraph("第三部分：基因变异及相应靶向/免疫药物解析")
    doc.sections[0].right_margin = Cm(2.0)
    second = doc.add_section(WD_SECTION.NEW_PAGE)
    second.right_margin = Cm(3.5)
    doc.add_paragraph("第四部分：附录")
    doc.save(docx_path)

    TemplateRenderer(log_level="ERROR")._normalize_final_section_layout(str(docx_path))

    margins = {round(section.right_margin.cm, 2) for section in Document(docx_path).sections}
    assert margins == {2.0}


def test_gene_list_table_is_compacted(tmp_path):
    docx_path = tmp_path / "gene_list.docx"
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Gene List for MLseq (n=358)"
    table.rows[1].cells[0].text = "ABL1"
    table.rows[1].cells[1].text = "ABL2"
    table.rows[2].cells[0].text = "ZRSR2"
    doc.save(docx_path)

    TemplateRenderer(log_level="ERROR")._compact_gene_list_tables(str(docx_path))

    doc = Document(docx_path)
    body_run = doc.tables[0].rows[1].cells[0].paragraphs[0].runs[0]
    assert body_run.font.size.pt == 6.0


def test_variant_table_layout_is_optimized(tmp_path):
    from docx.oxml.ns import qn

    docx_path = tmp_path / "variant_table.docx"
    doc = Document()
    table = doc.add_table(rows=3, cols=9)
    headers = [
        "基因名称", "基因突变信息", "基因突变信息", "基因突变信息",
        "基因突变信息", "基因突变信息", "基因突变信息",
        "靶向药物信息", "靶向药物信息",
    ]
    for idx, text in enumerate(headers):
        table.rows[0].cells[idx].text = text
    for idx in range(9):
        table.rows[1].cells[idx].text = "测试内容"
    values = [
        "TP53", "NM_000546.6", "17", "8", "c.817C>T,\np.R273C",
        "点突变", "76.12", "AZD1775（C）\nAlisertib（C）", "--",
    ]
    for idx, text in enumerate(values):
        table.rows[2].cells[idx].text = text
    doc.save(docx_path)

    TemplateRenderer(log_level="ERROR")._optimize_variant_table_layout(str(docx_path))

    doc = Document(docx_path)
    table = doc.tables[0]
    assert len(table.columns) == 4
    grid_widths = [
        int(col.get(qn("w:w")) or "0")
        for col in table._tbl.tblGrid.gridCol_lst
    ]
    assert grid_widths == [1100, 2600, 3000, 1600]
    assert "转录本：NM_000546.6" in table.rows[1].cells[1].text
    assert "频率：76.12%" in table.rows[1].cells[1].text
    assert table.rows[1].cells[0].paragraphs[0].runs[0].font.size.pt == 8.0


def test_real_no_variants_sample_regression_if_available():
    sample = (
        ROOT
        / "storage/uploads/2026-03-27/0ae95efb-606b-42f0-80d7-e21283c6415c"
        / "case3_no_variants_MLB0003.result.xlsx"
    )
    if not sample.exists():
        pytest.skip("real no-variants sample is not available in this checkout")

    excel_data = ExcelReader(
        config_dir=str(ROOT / "config"), log_level="ERROR"
    ).read(str(sample), include_tables=True)
    report_data = enhance_report_data(
        ReportData(),
        excel_data,
        base_path=str(ROOT),
    )

    assert report_data.get_field("total_variants_count") == 0
    assert report_data.get_table("variants") == []
    assert not any(
        (row.get("gene") or "").upper() == "ERBB2"
        for row in report_data.get_table("targeted_drug_tips")
    )
