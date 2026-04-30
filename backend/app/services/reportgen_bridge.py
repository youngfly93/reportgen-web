"""
Bridge service wrapping the upstream reportgen package.

Provides a clean API for the web layer to call ReportGenerator,
ExcelReader, ProjectDetector, and FieldMapper without knowing internals.
"""

import json
import sys
from pathlib import Path
from typing import Any, Optional

# Ensure upstream reportgen package is importable.
# Works in two layouts:
#   1. Vercel / self-contained: reportgen/ is in project root
#   2. Local dev: reportgen is in sibling directory ../基因组panel自动化系统/
from app.config import settings

_upstream = Path(str(settings.upstream_root))
if str(_upstream) not in sys.path:
    sys.path.insert(0, str(_upstream))

from reportgen.core.excel_reader import ExcelReader  # noqa: E402
from reportgen.core.field_mapper import FieldMapper  # noqa: E402
from reportgen.core.project_detector import ProjectDetector  # noqa: E402
from reportgen.core.report_generator import ReportGenerator  # noqa: E402
from reportgen.core.validation import validate_excel_data_common  # noqa: E402
from reportgen.models.excel_data import ExcelDataSource  # noqa: E402

DERIVED_REPORT_FIELDS = {
    "msi_status_cn",
    "msi",
    "msi_summary",
    "tmb",
    "tmb_status",
    "tmb_reference",
    "tmb_level_cn",
    "tmb_summary",
    "immuno_tips",
    "immuno_positive_genes",
    "immuno_negative_genes",
    "immuno_hyperprogression_genes",
}

EXCEL_OWNED_BIOMARKER_FIELDS = {
    "msi_status": ("MSI状态", "MSI结论", "MSI Status"),
    "msi_score": ("MSI百分比", "MSI评分", "MSI Score"),
    "tmb_value": ("TMB", "TMB值", "Tumor Mutation Burden"),
}


class ReportGenBridge:
    """Facade over the upstream reportgen library."""

    def __init__(self, config_dir: str, template_dir: str):
        self.config_dir = config_dir
        self.template_dir = template_dir
        self._generator: Optional[ReportGenerator] = None
        self._excel_reader: Optional[ExcelReader] = None
        self._detector: Optional[ProjectDetector] = None
        self._field_mapper: Optional[FieldMapper] = None

    @property
    def generator(self) -> ReportGenerator:
        if self._generator is None:
            self._generator = ReportGenerator(
                config_dir=self.config_dir,
                template_dir=self.template_dir,
                log_level="WARNING",
            )
        return self._generator

    @property
    def excel_reader(self) -> ExcelReader:
        if self._excel_reader is None:
            self._excel_reader = ExcelReader(
                config_dir=self.config_dir, log_level="WARNING"
            )
        return self._excel_reader

    @property
    def detector(self) -> ProjectDetector:
        if self._detector is None:
            self._detector = ProjectDetector(
                config_dir=self.config_dir, log_level="WARNING"
            )
        return self._detector

    @property
    def field_mapper(self) -> FieldMapper:
        if self._field_mapper is None:
            self._field_mapper = FieldMapper(
                config_dir=self.config_dir, log_level="WARNING"
            )
        return self._field_mapper

    def read_excel(self, excel_path: str) -> ExcelDataSource:
        """Read an Excel file and return structured data."""
        return self.excel_reader.read(excel_path)

    def get_sheet_names(self, excel_data: ExcelDataSource) -> list[str]:
        """Extract sheet names from parsed Excel data."""
        return list(excel_data.sheet_names) if excel_data.sheet_names else []

    def get_mapped_clinical_fields(self, excel_data: ExcelDataSource) -> dict[str, Any]:
        """
        Run FieldMapper on the Excel data to extract properly mapped clinical fields.

        This returns patient_name, sample_id, hospital, etc. — the real clinical
        fields — not the raw QC/technical data from excel_data.single_values.
        """
        try:
            report_data = self.field_mapper.map(excel_data)
            result = {}
            for k, v in report_data.context.items():
                # Skip table data (lists/dicts), keep only scalar clinical fields
                if isinstance(v, (list, dict)):
                    continue
                # Convert non-serializable types
                if hasattr(v, "item"):  # numpy scalar
                    v = v.item()
                if v is None:
                    continue
                if isinstance(v, float) and v != v:  # NaN
                    continue
                if v == "" or v == "-":
                    continue
                result[k] = v
            return result
        except Exception:
            # Fallback: return raw single_values if mapping fails
            return self._extract_raw_single_values(excel_data)

    def _extract_raw_single_values(self, excel_data: ExcelDataSource) -> dict[str, Any]:
        """Fallback: extract raw single values without field mapping."""
        result = {}
        for k, v in (excel_data.single_values or {}).items():
            if hasattr(v, "item"):
                v = v.item()
            if v is not None and not (isinstance(v, float) and v != v):
                result[k] = v
        return result

    def validate_excel_data(self, excel_data: ExcelDataSource) -> list[dict[str, str]]:
        """
        Validate Excel data for common issues and return warnings.

        Returns list of {level: "warning"|"error", message: str}
        """
        return validate_excel_data_common(excel_data)

    def get_table_data(
        self,
        excel_data: ExcelDataSource,
        table_name: str,
        page: int = 1,
        page_size: int = 50,
    ) -> dict[str, Any]:
        """Get table data for a specific sheet with pagination."""
        tables = excel_data.table_data or {}
        if table_name not in tables:
            return {
                "columns": [],
                "rows": [],
                "total_rows": 0,
                "page": page,
                "page_size": page_size,
            }

        df = tables[table_name]
        if hasattr(df, "to_dict"):
            # It's a DataFrame
            total = len(df)
            start = (page - 1) * page_size
            end = start + page_size
            page_df = df.iloc[start:end]
            columns = list(df.columns)
            rows = json.loads(
                page_df.to_json(
                    orient="records",
                    force_ascii=False,
                    default_handler=str,
                )
            )
        elif isinstance(df, list):
            total = len(df)
            start = (page - 1) * page_size
            end = start + page_size
            rows = df[start:end]
            columns = list(rows[0].keys()) if rows else []
        else:
            return {
                "columns": [],
                "rows": [],
                "total_rows": 0,
                "page": page,
                "page_size": page_size,
            }

        return {
            "columns": columns,
            "rows": rows,
            "total_rows": total,
            "page": page,
            "page_size": page_size,
        }

    def detect_project_type(
        self,
        excel_path: str,
        excel_data: Optional[ExcelDataSource] = None,
    ) -> dict[str, Any]:
        """
        Auto-detect project type from Excel filename + content.

        Args:
            excel_path: Path to Excel file (used for filename-based detection)
            excel_data: Optional parsed Excel data (used for content-based detection)
        """
        try:
            result = self.detector.detect(excel_path, excel_data=excel_data)
            if isinstance(result, dict):
                return {
                    "project_type": result.get("project_type"),
                    "project_name": result.get("project_name"),
                    "confidence": result.get("confidence"),
                    "detected": result.get("detected", False),
                }
            return {
                "project_type": None,
                "project_name": None,
                "confidence": None,
                "detected": False,
            }
        except Exception:
            return {
                "project_type": None,
                "project_name": None,
                "confidence": None,
                "detected": False,
            }

    def generate_report(
        self,
        excel_path: str,
        output_dir: str,
        template_name: Optional[str] = None,
        clinical_info: Optional[dict[str, Any]] = None,
        project_type: Optional[str] = None,
        project_name: Optional[str] = None,
        strict_mode: bool = False,
        template_contract_mode: str = "warn",
    ) -> dict[str, Any]:
        """
        Generate a single report.

        Returns dict with: success, output_file, duration, errors, warnings
        """
        # Resolve template
        if template_name:
            template_path = str(Path(self.template_dir) / template_name)
        else:
            template_path = str(
                Path(self.template_dir) / "aligned_template_with_cnv_fusion_hla_FIXED.docx"
            )

        # Read Excel first
        excel_data = self.read_excel(excel_path)

        # ========================================================
        # CRITICAL: Inject clinical_info so FieldMapper picks it up
        # ========================================================
        # FieldMapper iterates mapping.yaml fields and searches
        # excel_data.single_values for matching SYNONYMS (not var_names).
        # So we must inject values using the FIRST SYNONYM as key.
        if clinical_info:
            # Load mapping.yaml to get synonyms
            import yaml
            mapping_path = Path(self.config_dir) / "mapping.yaml"
            with open(mapping_path, "r", encoding="utf-8") as f:
                mapping_cfg = yaml.safe_load(f) or {}
            single_values_cfg = mapping_cfg.get("single_values", {})

            for var_name, value in clinical_info.items():
                if value is None or value == "":
                    continue
                if var_name in DERIVED_REPORT_FIELDS:
                    continue
                excel_owned_synonyms = EXCEL_OWNED_BIOMARKER_FIELDS.get(var_name)
                if excel_owned_synonyms and any(
                    key in excel_data.single_values for key in excel_owned_synonyms
                ):
                    # Prefer parsed Excel biomarkers over stale/hidden form data.
                    continue

                field_def = single_values_cfg.get(var_name)
                if field_def and isinstance(field_def, dict):
                    synonyms = field_def.get("synonyms", [])
                    if synonyms:
                        # Use first synonym as the key so FieldMapper's
                        # matches_column_name() will match it
                        excel_data.single_values[synonyms[0]] = value
                    else:
                        # Computed field (empty synonyms) — write to var_name directly
                        # These get picked up as fallback via report_data.get_field
                        excel_data.single_values[var_name] = value
                else:
                    # Unknown field — still inject under var_name for flexibility
                    excel_data.single_values[var_name] = value

            # Override filename-derived sample_id (the one used by patient_info.yaml lookup)
            # This prevents patient_info.yaml fallback from overwriting user form input
            if "sample_id" in clinical_info and clinical_info["sample_id"]:
                excel_data.metadata["sample_id_from_filename"] = clinical_info["sample_id"]

        # Resolve project_name from project_type if not provided. This must not
        # depend on fuzzy detection because operators can manually select a type
        # for files whose names do not contain "301/358".
        if project_type and not project_name:
            for entry in self.detector.project_types:
                if entry.get("id") == project_type:
                    project_name = entry.get("name")
                    break
            if not project_name:
                detect = self.detect_project_type(excel_path, excel_data=excel_data)
                if detect.get("detected"):
                    project_name = detect.get("project_name")

        result = self.generator.generate(
            excel_file=excel_path,
            template_file=template_path,
            output_dir=output_dir,
            excel_data=excel_data,
            strict_mode=strict_mode,
            return_context=False,
            template_contract_mode=template_contract_mode,
            project_type=project_type,
            project_name=project_name,
        )

        # Post-process: inject signature image if provided
        if result.get("success") and result.get("output_file") and clinical_info:
            sig_path = clinical_info.get("signature_image_path")
            if sig_path and Path(sig_path).exists():
                try:
                    self._insert_signature_image(result["output_file"], sig_path)
                except Exception as e:
                    if "warnings" not in result:
                        result["warnings"] = []
                    result["warnings"].append(f"签名图片插入失败: {e}")

        return result

    def _insert_signature_image(self, docx_path: str, signature_path: str) -> None:
        """Replace __SIG_IMG__ marker in the generated docx with an inline image."""
        from docx import Document as DocxDocument
        from docx.shared import Cm

        doc = DocxDocument(docx_path)
        MARKER = "__SIG_IMG__"

        for p in doc.paragraphs:
            if MARKER not in p.text:
                continue
            # Find run containing marker and split/replace
            for run in p.runs:
                if MARKER in run.text:
                    # Remove marker text, keep any prefix/suffix
                    parts = run.text.split(MARKER, 1)
                    run.text = parts[0]
                    # Insert image inline in the same run
                    run.add_picture(signature_path, width=Cm(2.5))
                    # Append suffix as new run if any
                    if len(parts) > 1 and parts[1]:
                        new_run = p.add_run(parts[1])
                        # Copy font from original run
                        if run.font.name:
                            new_run.font.name = run.font.name
                    break
            break

        doc.save(docx_path)
